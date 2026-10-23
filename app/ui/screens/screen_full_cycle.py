"""Экран «Полный цикл» (ТЗ §5.6, ШАГ 7): PDF → КМД, подготовка, сравнение, DEP-перенос.

Каждый этап — функции из src/ (импорт, без дублирования). Опциональный cleanup
промежуточных файлов — только удалением результатов, созданных этим циклом.
"""
from pathlib import Path

from PySide6.QtWidgets import QLabel, QComboBox, QPushButton, QCheckBox

from core.patches import set_prefilled
from core.settings import save_settings
from core.src_loader import load_src
from ui.screens.base_screen import BaseScreen
from ui.widgets import Card, FileSelect


class ScreenFullCycle(BaseScreen):
    action_label = "Запустить полный цикл"
    action_tip = "1) Парсинг PDF → 2) Подготовка → 3) Сравнение → 4) Перенос DEP (Ctrl+Enter)"

    def card_header(self) -> tuple[str, str]:
        return ("Полный цикл", "Все 4 этапа подряд с единым набором параметров")

    def build_form(self, card: Card):
        self.fs_pdf = FileSelect(".pdf", "Прошмандовка (PDF)")
        self.fs_xlsm = FileSelect(".xlsm", "Наша отчётная (.xlsm)")
        self.fs_xlsb = FileSelect(".xlsb", "DEP (.xlsb)")
        self.combo_sheet = QComboBox()      # лист для prepare_report
        self.combo_km = QComboBox()
        self.combo_dp = QComboBox()
        self.btn_refresh = QPushButton("Обновить списки (лист / КМ / DP)")
        self.btn_refresh.clicked.connect(self._load_lists)
        self.chk_cleanup = QCheckBox("Удалить промежуточные файлы после цикла")
        self.chk_cleanup.setChecked(bool(self.settings.get("cleanup_after_cycle", False)))

        for lbl, w in (("PDF (прошмандовка):", self.fs_pdf),
                       ("Отчётная (.xlsm):", self.fs_xlsm),
                       ("DEP (.xlsb):", self.fs_xlsb),
                       ("Лист отчётной (для группировки):", self.combo_sheet),
                       ("КМ:", self.combo_km), ("DP:", self.combo_dp),
                       ("", self.btn_refresh)):
            if lbl:
                card.add(QLabel(lbl))
            card.add(w)
        card.add(self.chk_cleanup)

        for key, fs in (("last_pdf", self.fs_pdf), ("last_xlsm", self.fs_xlsm),
                        ("last_xlsb", self.fs_xlsb)):
            v = self.settings.get(key, "")
            if v and Path(v).is_file():
                fs.set_text(v)

    def _load_lists(self):
        ok, err = self.fs_xlsm.validate()
        if not ok:
            self.log.append_line(f"⚠️ {err}")
            return
        try:
            from openpyxl import load_workbook
            import pandas as pd
            pr = load_src("prepare_report")
            p = Path(self.fs_xlsm.text())
            wb = load_workbook(p, read_only=True, data_only=True)
            sheets = list(wb.sheetnames)
            wb.close()
            sheet = self.combo_sheet.currentText() or (sheets[0] if sheets else "")
            raw = pd.read_excel(p, sheet_name=sheet, header=None,
                                engine="openpyxl", dtype=object)
            hr = pr.find_header_row(raw)
            if hr is None:
                raise RuntimeError("Строка шапки ('Марка' + 'Вес') не найдена")
            km_vals = sorted({pr.clean_str(v) for v in raw.iloc[hr + 1:, pr.COL_KM]
                              if pr.clean_str(v)})
            dp_vals = sorted({pr.clean_str(v) for v in raw.iloc[hr + 1:, pr.COL_DP]
                              if pr.clean_str(v)})
        except Exception as e:
            self.log.append_line(f"⚠️ {type(e).__name__}: {e}")
            return
        for combo, vals in ((self.combo_sheet, sheets), (self.combo_km, km_vals),
                            (self.combo_dp, dp_vals)):
            combo.blockSignals(True)
            combo.clear()
            combo.addItems([str(v) for v in vals])
            combo.blockSignals(False)
        self.log.append_line(f"📄 Листов: {len(sheets)}, КМ: {len(km_vals)}, DP: {len(dp_vals)}")

    def validate(self) -> str:
        for fs in (self.fs_pdf, self.fs_xlsm, self.fs_xlsb):
            ok, err = fs.validate()
            if not ok:
                return err
        for c, name in ((self.combo_sheet, "лист"), (self.combo_km, "КМ"),
                        (self.combo_dp, "DP")):
            if not c.currentText():
                return f"Не выбран {name} — нажмите «Обновить списки»"
        return ""

    def start_task(self):
        pdf = Path(self.fs_pdf.text())
        xlsm = Path(self.fs_xlsm.text())
        xlsb = Path(self.fs_xlsb.text())
        sheet, km, dp = (self.combo_sheet.currentText(), self.combo_km.currentText(),
                         self.combo_dp.currentText())
        cleanup = self.chk_cleanup.isChecked()

        def run():
            pk = load_src("pdf_to_kmd")
            pr = load_src("prepare_report")
            ck = load_src("compare_kmd")
            td = load_src("transfer_dep")
            mn = load_src("main")

            print("\n===== ЭТАП 1/4: ПАРСИНГ PDF =====")
            kmd_path = pk.parse_pdf(pdf)

            print("\n===== ЭТАП 2/4: ПОДГОТОВКА ОТЧЁТНОЙ =====")
            set_prefilled([("рабочий лист", sheet), ("Выберите КМ", km),
                           ("Выберите DP", dp)])
            report_path = pr.prepare_report(xlsm)

            print("\n===== ЭТАП 3/4: СРАВНЕНИЕ =====")
            comparison_path = ck.compare(kmd_path, report_path)

            print("\n===== ЭТАП 4/4: ПЕРЕНОС ИЗ DEP =====")
            # листы DEP и отчётной, КМ/DP — как в консоли: спросит gui_ask_choice
            # (через мост диалогов MainWindow); предзаполняем известные значения
            set_prefilled([("лист отчётной", sheet), ("Выберите КМ", km),
                           ("Выберите DP", dp)])
            wb, rep_sheet, _ = td.transfer_dep(xlsb, xlsm)
            final_path = mn.build_final_xlsx(
                xlsm_path=xlsm, main_sheet_name=rep_sheet, wb_from_transfer=wb,
                kmd_path=kmd_path, report_prepared_path=report_path,
                comparison_path=comparison_path)

            if cleanup:
                print("\n--- CLEANUP ---")
                for tmp in (kmd_path, report_path, comparison_path):
                    try:
                        Path(tmp).unlink(missing_ok=True)
                        print(f"🗑 Удалён: {Path(tmp).name}")
                    except OSError as e:
                        print(f"⚠️ Не удалён {Path(tmp).name}: {e}")
            return final_path

        return (run, (), {}, None)

    def on_success(self, result):
        self.settings.update(last_pdf=self.fs_pdf.text(),
                             last_xlsm=self.fs_xlsm.text(),
                             last_xlsb=self.fs_xlsb.text(),
                             last_sheet=self.combo_sheet.currentText(),
                             last_km=self.combo_km.currentText(),
                             last_dp=self.combo_dp.currentText(),
                             cleanup_after_cycle=self.chk_cleanup.isChecked())
        save_settings(self.settings)
        self.show_result_button(Path(result))
